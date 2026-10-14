"""営業リストとフォームのスクリーンショットから、入力内容の確認ページ(HTML)を作る

  python3 tools/build_review.py <スクリーンショットのあるフォルダ> <出力.html> <config.json> <営業リスト.xlsx> <最初のNo> [手作業扱い.json] [見出し]

- <スクリーンショットのあるフォルダ> には screenshots/ が入っていること（src/index.js の出力先）
- 手作業扱い.json: {"No": "理由"} 。ツールが欄を読み違えた会社などを「手作業が必要」に分ける
- SA登録の列に記載がある会社（業者マスタに登録あり・要確認）は載せない
"""
import base64, html, io, json, sys
from pathlib import Path
import openpyxl
from PIL import Image

run = Path(sys.argv[1])
out = Path(sys.argv[2])
cfg = json.load(open(sys.argv[3]))

wb = openpyxl.load_workbook(sys.argv[4])
MIN = int(sys.argv[5])
EXTRA = json.load(open(sys.argv[6])) if len(sys.argv) > 6 else {}
rows = [r for r in list(wb["自動チェック詳細"].iter_rows(values_only=True))[1:] if r[0] and int(r[0]) >= MIN]
main = {str(r[0]): r for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True) if r[0]}
rows.sort(key=lambda r: int(r[0]))



def img_uri(no):
    shots = sorted((run / "screenshots").glob(f"{int(no):02d}_*_form.png")) or sorted((run / "screenshots").glob(f"{int(no):02d}_*.png"))
    if not shots:
        return None
    im = Image.open(shots[0]).convert("RGB")
    if im.width > 640:
        im = im.resize((640, round(im.height * 640 / im.width)))
    if im.height > 1800:
        im = im.crop((0, 0, im.width, 1800))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=62, optimize=True)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


def esc(s):
    return html.escape(str(s or ""))


cards = []
n_ok = n_manual = 0
for r in rows:
    no, company, url, status, _, captcha, filled, missing = r[:8]
    if not str(status).startswith("入力済"):
        continue
    if main.get(str(no)) and main[str(no)][14]:  # SA登録の列に記載がある会社（登録あり・要確認）は載せない
        continue
    # 目に見えるCAPTCHA・必須の空き・文字数オーバーがある会社は、送信時に手作業が要る
    manual = str(no) in EXTRA or bool(missing) or "文字数オーバー" in str(status) or "自動送信不可" in str(status) or bool(captcha and "不可視" not in str(captcha))
    n_manual += manual
    n_ok += not manual
    pairs = []
    for line in str(filled or "").split("\n"):
        if not line.strip():
            continue
        k, _, v = line.partition("：")
        pairs.append(f"<tr><th scope='row'>{esc(k)}</th><td>{esc(v)}</td></tr>")
    note = []
    soft = ""
    if captcha and "不可視" in str(captcha):
        soft = "<p class='soft'>目に見えない認証（reCAPTCHA）あり。通常は手作業なしで送れます</p>"
    elif captcha:
        note.append(esc(captcha))
    if missing:
        note.append("未入力の必須：" + esc(missing))
    if str(no) in EXTRA:
        note.append(esc(EXTRA[str(no)]))
    m = main.get(str(no))
    if m:
        memo = [l for l in str(m[16] or "").split(chr(10)) if l.startswith("[調査メモ]")]
        soft += f"<p class='soft'>{esc(m[2])}｜{esc(m[3])}｜規模：{esc(m[4])}</p>" + "".join(f"<p class='soft'>{esc(l)}</p>" for l in memo)
    uri = img_uri(no)
    shot = f"<img src='{uri}' alt='{esc(company)}の入力後のフォーム' loading='lazy'>" if uri else "<p class='noshot'>フォーム部分の画像を撮れませんでした</p>"
    cards.append(f"""
<article class="card" data-kind="{'manual' if manual else 'ok'}" id="no{no}">
  <header>
    <span class="no">No.{esc(no)}</span>
    <h2>{esc(company)}</h2>
    <span class="pill {'warn' if manual else 'good'}">{'手作業が必要' if manual else 'そのまま送れそう'}</span>
  </header>
  <a class="url" href="{esc(url)}" target="_blank" rel="noopener">{esc(url)}</a>
  {'<p class="note">' + ' ／ '.join(note) + '</p>' if note else ''}{soft}
  <div class="body">
    <div class="pairs"><table><caption>入力した内容</caption><tbody>{''.join(pairs)}</tbody></table></div>
    <figure>{shot}</figure>
  </div>
</article>""")


def render(t):
    s = cfg["sender"]
    return (t.replace("{{自分の名前}}", s["lastName"] + s["firstName"]).replace("{{自社名}}", s["company"])
            .replace("{{部署}}", s["department"]).replace("{{電話}}", s["tel"]).replace("{{メール}}", s["email"])
            .replace("{{企業名}}", "（各社の会社名）"))


HEADING = sys.argv[7] if len(sys.argv) > 7 else f"追加リスト（No.{MIN}〜）フォーム入力確認"
page = open(Path(__file__).with_name("review_template.html"), encoding="utf-8").read().replace("%%TITLE%%", "フォーム入力確認").replace("%%HEADING%%", HEADING)
page = (page.replace("%%CARDS%%", "\n".join(cards))
        .replace("%%N_ALL%%", str(n_ok + n_manual)).replace("%%N_OK%%", str(n_ok)).replace("%%N_MANUAL%%", str(n_manual))
        .replace("%%SUBJECT%%", esc(cfg["subject"]))
        .replace("%%MSG_FULL%%", esc(render(cfg["message"]))).replace("%%MSG_SHORT%%", esc(render(cfg["messageShort"])))
        .replace("%%LEN_FULL%%", str(len(render(cfg["message"])))).replace("%%LEN_SHORT%%", str(len(render(cfg["messageShort"])))))
out.write_text(page, encoding="utf-8")
print(out, round(out.stat().st_size / 1e6, 1), "MB", n_ok, n_manual)
